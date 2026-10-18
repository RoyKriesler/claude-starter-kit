#!/usr/bin/env python
"""Render-and-check gate for file deliverables. Used by the builder agent before it may say "done".

"Written and rendered-correct are different claims." This script proves the first half mechanically:
the file opens in the real application, renders, and carries no machine-detectable defect. The
second half (does the rendered thing look right, say the right thing) is the caller's eyes on the
PNGs this script writes. A PASS here is necessary, never sufficient.

Usage:
  python render_check.py xlsx    FILE [--expect Sheet!A1=10,B2=abc] [--require-formulas]
  python render_check.py docx    FILE [--rtl]
  python render_check.py pptx    FILE
  python render_check.py pdf     FILE [--expect-a4] [--expect-font Heebo]
  python render_check.py html    PATH_OR_URL [--viewport 1280x800] [--mobile]
  python render_check.py dataset FILE [any validate_dataset.py flags]
  python render_check.py hebrew  FILE          (docx / pptx / xlsx / pdf / html / md / txt)
  common: [--out DIR] [--dpi N] [--max-pages N] [--timeout SEC] [--json]

Writes artefacts (pdf, png per page or slide, screenshot) under --out (default
%TEMP%/render_check/<stem>-<HHMMSS>/) and prints a JSON report. Exit 0 = PASS, 1 = FAIL, 2 = usage.

Machine notes (the Windows box, 2026-09-04): no LibreOffice, no pdftoppm. Office 16 COM is the
renderer for xlsx / docx / pptx; PyMuPDF rasterises PDFs; Playwright chromium renders HTML. Every
COM job runs in a child process with a timeout; on timeout the child and any Office process it
started are killed, and the gate reports FAIL with reason "timeout" instead of hanging the session.
OneDrive paths hang Word COM, so inputs are copied to a local temp dir first.
"""
import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
HEBREW = re.compile(r"[֐-׿]")
LATIN = re.compile(r"[A-Za-z]")
BRACKETS = re.compile(r"[()\[\]{}/|<>]")
OFFICE_IMAGES = {"xlsx": "EXCEL.EXE", "docx": "WINWORD.EXE", "pptx": "POWERPNT.EXE"}


# ----------------------------------------------------------------------------- helpers

def out_dir(args, path):
    if args.out:
        d = args.out
    else:
        stem = os.path.splitext(os.path.basename(path))[0] or "site"
        d = os.path.join(tempfile.gettempdir(), "render_check", f"{stem}-{time.strftime('%H%M%S')}")
    os.makedirs(d, exist_ok=True)
    return os.path.abspath(d)


def local_copy(path, d):
    """Copy the input next to the artefacts so Office never sees a OneDrive path."""
    dst = os.path.join(d, "input" + os.path.splitext(path)[1].lower())
    shutil.copy2(path, dst)
    return os.path.normpath(dst)


def office_pids(image):
    try:
        txt = subprocess.run(["tasklist", "/FI", f"IMAGENAME eq {image}", "/FO", "CSV", "/NH"],
                             capture_output=True, text=True, timeout=20).stdout
    except Exception:
        return set()
    return {int(r.split('","')[1]) for r in txt.splitlines() if r.startswith('"')}


def kill_pids(pids):
    for p in pids:
        subprocess.run(["taskkill", "/PID", str(p), "/F", "/T"], capture_output=True)


def pid_alive(pid):
    try:
        txt = subprocess.run(["tasklist", "/FI", f"PID eq {pid}", "/FO", "CSV", "/NH"], capture_output=True, text=True, timeout=20).stdout
    except Exception:
        return False
    return any(r.startswith('"') for r in txt.splitlines())


def read_pidfile(path):
    try:
        return {int(x) for x in open(path).read().split() if x.strip().isdigit()}
    except Exception:
        return set()


def run_worker(kind, payload, timeout):
    """Run the COM part in a child python with a hard timeout. Returns dict.

    The worker writes the pid of the Office instance it started to payload["pidfile"] right after
    DispatchEx, so this parent only ever kills that instance: never another runner's Word or Excel
    that happens to be open at the same time (parallel eval waves, the user's own documents)."""
    image = OFFICE_IMAGES[kind]
    payload = dict(payload)
    payload["pidfile"] = os.path.join(os.path.dirname(payload["path"]), f"office-{kind}.pid")
    cmd = [sys.executable, os.path.abspath(__file__), "--worker", kind, json.dumps(payload)]
    t0 = time.time()
    try:
        cp = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout, encoding="utf-8")
    except subprocess.TimeoutExpired:
        mine = read_pidfile(payload["pidfile"])
        kill_pids(mine)
        return {"ok": False, "error": f"timeout after {timeout}s in {image} COM; killed own instance {sorted(mine) or '(pid unknown)'}", "elapsed": round(time.time() - t0, 1)}
    leftover = {p for p in read_pidfile(payload["pidfile"]) if pid_alive(p)}
    if leftover:
        kill_pids(leftover)
    if cp.returncode != 0:
        err = (cp.stderr or cp.stdout or "").strip().splitlines()
        return {"ok": False, "error": " | ".join(err[-4:]) or f"worker exit {cp.returncode}", "elapsed": round(time.time() - t0, 1)}
    try:
        res = json.loads(cp.stdout.strip().splitlines()[-1])
    except Exception:
        return {"ok": False, "error": f"worker output unreadable: {cp.stdout[-300:]}", "elapsed": round(time.time() - t0, 1)}
    res["elapsed"] = round(time.time() - t0, 1)
    if leftover:
        res.setdefault("warnings", []).append(f"{image} still alive 5 s after Quit; killed own instance {sorted(leftover)}")
    return res


def rasterise(pdf_path, d, dpi, max_pages):
    import fitz
    doc = fitz.open(pdf_path)
    pngs, text_chars, sizes, fonts, blank = [], 0, [], set(), []
    for i, page in enumerate(doc):
        w_mm, h_mm = round(page.rect.width * 25.4 / 72, 1), round(page.rect.height * 25.4 / 72, 1)
        sizes.append([w_mm, h_mm])
        txt = page.get_text()
        text_chars += len(txt.strip())
        if not txt.strip() and not page.get_images():
            blank.append(i + 1)
        for f in page.get_fonts():
            fonts.add(f[3].split("+")[-1])
        if i < max_pages:
            p = os.path.join(d, f"page-{i + 1:02d}.png")
            page.get_pixmap(dpi=dpi).save(p)
            pngs.append(p)
    return {"pages": doc.page_count, "png": pngs, "text_chars": text_chars, "page_size_mm": sizes,
            "fonts": sorted(fonts), "blank_pages": blank}


def extract_text(path):
    """Text lines from any supported deliverable, for the hebrew scan. Returns list of (where, line)."""
    ext = os.path.splitext(path)[1].lower()
    lines = []
    if ext == ".docx":
        import docx
        d = docx.Document(path)
        for i, p in enumerate(d.paragraphs, 1):
            if p.text.strip():
                lines.append((f"p{i}", p.text))
        for ti, t in enumerate(d.tables, 1):
            for ri, row in enumerate(t.rows, 1):
                cells = [c.text for c in row.cells]
                lines.append((f"table{ti}.row{ri}", " | ".join(cells)))
                if any(HEBREW.search(c) for c in cells):
                    lines.append((f"table{ti}.row{ri}", "__TABLE_WITH_HEBREW__"))
    elif ext == ".pptx":
        from pptx import Presentation
        for si, s in enumerate(Presentation(path).slides, 1):
            for sh in s.shapes:
                if sh.has_text_frame:
                    for pi, p in enumerate(sh.text_frame.paragraphs, 1):
                        if p.text.strip():
                            lines.append((f"slide{si}.{sh.name}.p{pi}", p.text))
                if getattr(sh, "has_chart", False) and sh.has_chart:
                    ch = sh.chart
                    if ch.has_title:
                        lines.append((f"slide{si}.{sh.name}.chart.title", ch.chart_title.text_frame.text))
                    try:
                        for c in ch.plots[0].categories:
                            lines.append((f"slide{si}.{sh.name}.chart.category", str(c)))
                    except Exception:
                        pass
    elif ext in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True)
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.strip():
                        lines.append((f"{ws.title}!{c.coordinate}", c.value))
    elif ext == ".pdf":
        import fitz
        for i, page in enumerate(fitz.open(path), 1):
            for li, l in enumerate(page.get_text().splitlines(), 1):
                if l.strip():
                    lines.append((f"page{i}.l{li}", l))
    elif ext in (".html", ".htm"):
        raw = open(path, encoding="utf-8", errors="replace").read()
        raw = re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", raw, flags=re.S | re.I)
        raw = re.sub(r"<br\s*/?>|</(p|div|li|h\d|tr|td|th)>", "\n", raw, flags=re.I)
        raw = re.sub(r"<[^>]+>", " ", raw)
        for li, l in enumerate(raw.splitlines(), 1):
            if l.strip():
                lines.append((f"l{li}", re.sub(r"\s+", " ", l).strip()))
    else:  # md, txt
        fenced, in_table = False, False
        for li, l in enumerate(open(path, encoding="utf-8", errors="replace").read().splitlines(), 1):
            if l.strip().startswith("```"):
                fenced = not fenced
                continue
            if fenced or not l.strip():
                continue
            if l.lstrip().startswith("|") and HEBREW.search(l):
                lines.append((f"l{li}", "__TABLE_WITH_HEBREW__"))
            lines.append((f"l{li}", l))
    return lines


def hebrew_scan(lines):
    """the Hebrew language rule, mechanical part. Returns findings; caller reads each before judging."""
    findings = []
    for where, l in lines:
        if l == "__TABLE_WITH_HEBREW__":
            findings.append({"where": where, "rule": "table-with-hebrew", "token": "(table)"})
            continue
        if not HEBREW.search(l):
            continue
        # Latin inside a Hebrew line: report each Latin run
        for m in re.finditer(r"[A-Za-z][A-Za-z0-9_.@:+\-]*", l):
            findings.append({"where": where, "rule": "latin-in-hebrew-line", "token": m.group(0)})
        for m in BRACKETS.finditer(l):
            findings.append({"where": where, "rule": "bracket-in-hebrew-line", "token": m.group(0)})
    return findings


def finish(report, as_json=True):
    report["status"] = "FAIL" if report.get("fail") else "PASS"
    report.pop("fail", None)
    print(json.dumps(report, ensure_ascii=False, indent=1))
    return 0 if report["status"] == "PASS" else 1


# ----------------------------------------------------------------------------- COM workers (child process)

def _pid_of_hwnd(hwnd):
    import win32process
    return win32process.GetWindowThreadProcessId(int(hwnd))[1]


def _note_pid(p, pids):
    try:
        with open(p["pidfile"], "w") as f:
            f.write(" ".join(str(x) for x in sorted(pids)))
    except Exception:
        pass


def _wait_exit(pids, seconds=5):
    t0 = time.time()
    while time.time() - t0 < seconds and any(pid_alive(x) for x in pids):
        time.sleep(0.25)


def _com_error_text(e):
    """pywintypes.com_error carries the application's message in excepinfo[2]; surface it."""
    try:
        info = e.excepinfo
        if info and info[2]:
            return f"{type(e).__name__}: {info[2].strip()}"
    except Exception:
        pass
    return f"{type(e).__name__}: {e}"


def _run_com(fn, p):
    try:
        return fn(p)
    except Exception as e:
        return {"ok": False, "error": _com_error_text(e)}


def worker_xlsx(p):
    import win32com.client
    xl = win32com.client.DispatchEx("Excel.Application")
    h = xl.Hwnd
    pids = {_pid_of_hwnd(h() if callable(h) else h)}
    _note_pid(p, pids)
    xl.Visible = False
    xl.DisplayAlerts = False
    res = {"ok": True, "errors": [], "formulas": 0, "sheets": [], "expect": []}
    try:
        wb = xl.Workbooks.Open(p["path"], UpdateLinks=0, ReadOnly=True)
        xl.CalculateFull()
        for sh in wb.Worksheets:
            n = 0
            ur = sh.UsedRange
            for cell in ur.Cells:
                if cell.HasFormula:
                    res["formulas"] += 1
                    n += 1
                    txt = str(cell.Text)
                    if txt.startswith("#"):
                        res["errors"].append({"cell": f"{sh.Name}!{cell.Address}", "formula": str(cell.Formula), "shows": txt})
            res["sheets"].append({"name": sh.Name, "used": str(ur.Address), "formulas": n})
        for e in p.get("expect", []):
            ref, want = e.split("=", 1)
            sh_name, addr = (ref.split("!", 1) if "!" in ref else (wb.Worksheets(1).Name, ref))
            got = wb.Worksheets(sh_name).Range(addr).Value
            got_s = str(got).rstrip("0").rstrip(".") if isinstance(got, float) else str(got)
            res["expect"].append({"cell": ref, "want": want, "got": str(got), "ok": got_s == want.strip()})
        wb.Close(False)
    finally:
        xl.Quit()
        del xl
        _wait_exit(pids)
    return res


def worker_docx(p):
    import win32com.client
    before = office_pids("WINWORD.EXE")
    w = win32com.client.DispatchEx("Word.Application")
    pids = office_pids("WINWORD.EXE") - before
    _note_pid(p, pids)
    w.Visible = False
    w.DisplayAlerts = 0
    res = {"ok": True}
    try:
        doc = w.Documents.Open(p["path"], ReadOnly=True, AddToRecentFiles=False)
        res["word_paragraphs"] = doc.Paragraphs.Count
        if p.get("rtl"):
            ltr = []
            for i in range(1, min(doc.Paragraphs.Count, 400) + 1):
                para = doc.Paragraphs(i)
                if HEBREW.search(para.Range.Text or "") and para.Format.Alignment != 2:
                    ltr.append({"paragraph": i, "alignment": int(para.Format.Alignment), "text": (para.Range.Text or "")[:60]})
            res["hebrew_paragraphs_not_right_aligned"] = ltr
        doc.SaveAs2(p["pdf"], FileFormat=17)
        doc.Close(False)
    finally:
        w.Quit()
        del w
        _wait_exit(pids)
    return res


def worker_pptx(p):
    import win32com.client
    before = office_pids("POWERPNT.EXE")
    app = win32com.client.DispatchEx("PowerPoint.Application")
    pids = office_pids("POWERPNT.EXE") - before      # HWND is not reachable via dynamic dispatch
    _note_pid(p, pids)
    res = {"ok": True}
    try:
        pres = app.Presentations.Open(p["path"], ReadOnly=True, Untitled=False, WithWindow=False)
        res["slides"] = pres.Slides.Count
        os.makedirs(p["png_dir"], exist_ok=True)
        pres.Export(p["png_dir"], "PNG", p.get("w", 1280), p.get("h", 720))
        pres.Close()
    finally:
        app.Quit()
        del app
        _wait_exit(pids)
    return res


WORKERS = {"xlsx": worker_xlsx, "docx": worker_docx, "pptx": worker_pptx}


# ----------------------------------------------------------------------------- commands

def cmd_xlsx(args):
    d = out_dir(args, args.file)
    rep = {"kind": "xlsx", "file": os.path.abspath(args.file), "out": d, "findings": []}
    expect = [e for e in (args.expect or "").split(",") if e.strip()]
    r = run_worker("xlsx", {"path": local_copy(args.file, d), "expect": expect}, args.timeout)
    rep["elapsed_s"] = r.get("elapsed")
    if not r.get("ok"):
        rep["findings"].append({"severity": "FAIL", "check": "opens-and-recalculates", "detail": r.get("error")})
        rep["fail"] = True
        return finish(rep)
    rep["stats"] = {"formulas": r["formulas"], "sheets": r["sheets"]}
    for e in r["errors"]:
        rep["findings"].append({"severity": "FAIL", "check": "formula-error-value", "detail": f"{e['cell']} {e['formula']} shows {e['shows']}"})
    if args.require_formulas and r["formulas"] == 0:
        rep["findings"].append({"severity": "FAIL", "check": "has-formulas", "detail": "no formula cells; totals are hard-coded values"})
    for e in r["expect"]:
        rep["findings"].append({"severity": "PASS" if e["ok"] else "FAIL", "check": "expected-value", "detail": f"{e['cell']} want {e['want']} got {e['got']}"})
    for wmsg in r.get("warnings", []):
        rep["findings"].append({"severity": "WARN", "check": "process", "detail": wmsg})
    rep["fail"] = any(f["severity"] == "FAIL" for f in rep["findings"])
    return finish(rep)


def cmd_docx(args):
    d = out_dir(args, args.file)
    rep = {"kind": "docx", "file": os.path.abspath(args.file), "out": d, "findings": []}
    try:
        with zipfile.ZipFile(args.file) as z:
            xml = z.read("word/document.xml").decode("utf-8", "replace")
        rep["stats"] = {"w_p": xml.count("<w:p>") + xml.count("<w:p "), "w_jc": xml.count("<w:jc "), "w_bidi": xml.count("<w:bidi")}
    except Exception as e:
        rep["findings"].append({"severity": "FAIL", "check": "zip-structure", "detail": f"{type(e).__name__}: {e}"})
        rep["fail"] = True
        return finish(rep)
    pdf = os.path.join(d, "render.pdf")
    r = run_worker("docx", {"path": local_copy(args.file, d), "pdf": pdf, "rtl": bool(args.rtl)}, args.timeout)
    rep["elapsed_s"] = r.get("elapsed")
    if not r.get("ok") or not os.path.exists(pdf):
        rep["findings"].append({"severity": "FAIL", "check": "word-opens-and-exports-pdf", "detail": r.get("error", "no pdf written")})
        rep["fail"] = True
        return finish(rep)
    ras = rasterise(pdf, d, args.dpi, args.max_pages)
    rep["artifacts"] = {"pdf": pdf, "png": ras["png"]}
    rep["stats"].update({k: ras[k] for k in ("pages", "text_chars", "page_size_mm", "fonts", "blank_pages")})
    rep["stats"]["word_paragraphs"] = r.get("word_paragraphs")
    if ras["pages"] == 0 or ras["text_chars"] == 0:
        rep["findings"].append({"severity": "FAIL", "check": "renders-with-text", "detail": f"pages={ras['pages']} text_chars={ras['text_chars']}"})
    if ras["blank_pages"]:
        rep["findings"].append({"severity": "WARN", "check": "blank-pages", "detail": str(ras["blank_pages"])})
    if args.rtl:
        ltr = r.get("hebrew_paragraphs_not_right_aligned", [])
        if ltr:
            rep["findings"].append({"severity": "FAIL", "check": "rtl-alignment", "detail": f"{len(ltr)} Hebrew paragraphs not right-aligned in Word: " + "; ".join(f"p{x['paragraph']} align={x['alignment']}" for x in ltr[:8])})
        if rep["stats"]["w_jc"]:
            rep["findings"].append({"severity": "WARN", "check": "rtl-jc-elements", "detail": f"{rep['stats']['w_jc']} <w:jc> elements; hebrew-office says 0 for RTL paragraphs"})
    for wmsg in r.get("warnings", []):
        rep["findings"].append({"severity": "WARN", "check": "process", "detail": wmsg})
    rep["fail"] = any(f["severity"] == "FAIL" for f in rep["findings"])
    return finish(rep)


def cmd_pptx(args):
    d = out_dir(args, args.file)
    rep = {"kind": "pptx", "file": os.path.abspath(args.file), "out": d, "findings": []}
    try:
        from pptx import Presentation
        pres = Presentation(args.file)
        slide_text = []
        for si, s in enumerate(pres.slides, 1):
            t = " ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame).strip()
            slide_text.append(len(t))
        rep["stats"] = {"slides_pptx": len(slide_text), "text_chars_per_slide": slide_text}
    except Exception as e:
        rep["findings"].append({"severity": "FAIL", "check": "zip-structure", "detail": f"{type(e).__name__}: {e}"})
        rep["fail"] = True
        return finish(rep)
    png_dir = os.path.join(d, "slides")
    r = run_worker("pptx", {"path": local_copy(args.file, d), "png_dir": png_dir}, args.timeout)
    rep["elapsed_s"] = r.get("elapsed")
    pngs = sorted(os.path.join(png_dir, f) for f in os.listdir(png_dir)) if os.path.isdir(png_dir) else []
    if not r.get("ok") or not pngs:
        rep["findings"].append({"severity": "FAIL", "check": "powerpoint-opens-and-exports-png", "detail": r.get("error", "no png written")})
        rep["fail"] = True
        return finish(rep)
    rep["artifacts"] = {"png": pngs[: args.max_pages]}
    rep["stats"]["slides_exported"] = len(pngs)
    empty = [i + 1 for i, n in enumerate(slide_text) if n == 0]
    if empty:
        rep["findings"].append({"severity": "WARN", "check": "slides-without-text", "detail": str(empty)})
    # geometry: a shape whose box crosses the slide boundary (charts cut at the edge; eval case 8, 2026-09-04)
    sw, sh_ = pres.slide_width, pres.slide_height
    outside = []
    for si, s in enumerate(pres.slides, 1):
        for shp in s.shapes:
            if None in (shp.left, shp.top, shp.width, shp.height):
                continue
            over_x = (shp.left + shp.width) - sw
            over_y = (shp.top + shp.height) - sh_
            if shp.left < 0 or shp.top < 0 or over_x > 0 or over_y > 0:
                outside.append(f"slide{si}.{shp.name} ({'chart' if getattr(shp, 'has_chart', False) and shp.has_chart else 'shape'}) over by x={max(over_x, 0) / 914400:.2f}in y={max(over_y, 0) / 914400:.2f}in")
    if outside:
        rep["findings"].append({"severity": "WARN", "check": "shape-outside-slide", "detail": "; ".join(outside[:8])})
    for where, l in extract_text(args.file):
        if "chart" in where and HEBREW.search(l) and "(" in l:
            rep["findings"].append({"severity": "FAIL", "check": "hebrew-chart-text-with-parens", "detail": f"{where}: {l[:60]} (renders scrambled, hebrew-office §2)"})
    for wmsg in r.get("warnings", []):
        rep["findings"].append({"severity": "WARN", "check": "process", "detail": wmsg})
    rep["fail"] = any(f["severity"] == "FAIL" for f in rep["findings"])
    return finish(rep)


def cmd_pdf(args):
    d = out_dir(args, args.file)
    rep = {"kind": "pdf", "file": os.path.abspath(args.file), "out": d, "findings": []}
    try:
        ras = rasterise(args.file, d, args.dpi, args.max_pages)
    except Exception as e:
        rep["findings"].append({"severity": "FAIL", "check": "pdf-opens", "detail": f"{type(e).__name__}: {e}"})
        rep["fail"] = True
        return finish(rep)
    rep["artifacts"] = {"png": ras["png"]}
    rep["stats"] = {k: ras[k] for k in ("pages", "text_chars", "page_size_mm", "fonts", "blank_pages")}
    if ras["pages"] == 0 or ras["text_chars"] == 0:
        rep["findings"].append({"severity": "FAIL", "check": "renders-with-text", "detail": f"pages={ras['pages']} text_chars={ras['text_chars']}"})
    if ras["blank_pages"]:
        rep["findings"].append({"severity": "WARN", "check": "blank-pages", "detail": str(ras["blank_pages"])})
    if args.expect_a4:
        bad = [s for s in ras["page_size_mm"] if abs(s[0] - 210) > 1.5 or abs(s[1] - 297) > 1.5]
        rep["findings"].append({"severity": "FAIL" if bad else "PASS", "check": "a4-page-size", "detail": f"sizes mm {ras['page_size_mm'][:3]}" + (" (Letter = 215.9x279.4: missing @page rule)" if bad else "")})
    if args.expect_font:
        hit = [f for f in ras["fonts"] if args.expect_font.lower() in f.lower()]
        rep["findings"].append({"severity": "PASS" if hit else "FAIL", "check": "font-embedded", "detail": f"want {args.expect_font}, fonts in pdf: {ras['fonts']}"})
    rep["fail"] = any(f["severity"] == "FAIL" for f in rep["findings"])
    return finish(rep)


def cmd_html(args):
    from playwright.sync_api import sync_playwright
    target = args.file
    is_url = target.startswith(("http://", "https://"))
    d = out_dir(args, "site" if is_url else target)
    rep = {"kind": "html", "file": target if is_url else os.path.abspath(target), "out": d, "findings": [], "console": []}
    if not is_url:
        if not os.path.exists(target):
            rep["findings"].append({"severity": "FAIL", "check": "file-exists", "detail": target})
            rep["fail"] = True
            return finish(rep)
        target = "file:///" + os.path.abspath(target).replace(os.sep, "/")
    w, h = [int(x) for x in args.viewport.lower().split("x")]
    shots = []
    with sync_playwright() as p:
        b = p.chromium.launch()
        for vw, vh, tag in ([(w, h, "desktop")] + ([(390, 844, "mobile")] if args.mobile else [])):
            pg = b.new_page(viewport={"width": vw, "height": vh})
            logs = []
            pg.on("console", lambda m: logs.append({"type": m.type, "text": m.text}) if m.type in ("error", "warning") else None)
            pg.on("pageerror", lambda e: logs.append({"type": "pageerror", "text": str(e)}))
            pg.on("requestfailed", lambda r: logs.append({"type": "requestfailed", "text": f"{r.url} {r.failure}"}))
            try:
                resp = pg.goto(target, wait_until="networkidle", timeout=args.timeout * 1000)
                status = resp.status if resp else None
            except Exception as e:
                rep["findings"].append({"severity": "FAIL", "check": f"loads-{tag}", "detail": f"{type(e).__name__}: {str(e)[:200]}"})
                pg.close()
                continue
            pg.wait_for_timeout(500)
            body = pg.inner_text("body") if pg.locator("body").count() else ""
            shot = os.path.join(d, f"screenshot-{tag}.png")
            pg.screenshot(path=shot, full_page=True)
            shots.append(shot)
            overflow = pg.evaluate("() => document.documentElement.scrollWidth > document.documentElement.clientWidth + 1")
            broken_imgs = pg.evaluate("() => Array.from(document.images).filter(i => i.complete && i.naturalWidth === 0).map(i => i.getAttribute('src'))")
            rep.setdefault("stats", {})[tag] = {"http_status": status, "title": pg.title(), "body_text_chars": len(body.strip()),
                                                 "horizontal_overflow": overflow, "broken_images": broken_imgs}
            rep["console"].extend({"viewport": tag, **l} for l in logs)
            if status and status >= 400:
                rep["findings"].append({"severity": "FAIL", "check": f"http-status-{tag}", "detail": str(status)})
            if len(body.strip()) == 0:
                rep["findings"].append({"severity": "FAIL", "check": f"visible-text-{tag}", "detail": "body inner_text empty"})
            errs = [l for l in logs if l["type"] in ("error", "pageerror", "requestfailed")]
            if errs:
                rep["findings"].append({"severity": "FAIL", "check": f"zero-console-errors-{tag}", "detail": f"{len(errs)}: " + "; ".join(e["text"][:120] for e in errs[:5])})
            if broken_imgs:
                rep["findings"].append({"severity": "FAIL", "check": f"images-load-{tag}", "detail": str(broken_imgs[:5])})
            if overflow:
                rep["findings"].append({"severity": "WARN", "check": f"horizontal-overflow-{tag}", "detail": "page wider than viewport"})
            pg.close()
        b.close()
    rep["artifacts"] = {"png": shots}
    rep["fail"] = any(f["severity"] == "FAIL" for f in rep["findings"])
    return finish(rep)


def cmd_dataset(args, passthrough):
    script = os.path.join(HERE, "validate_dataset.py")
    cmd = [sys.executable, script, args.file] + [a for a in passthrough if a != "--json"] + ["--json"]
    cp = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8")
    rep = {"kind": "dataset", "file": os.path.abspath(args.file), "command": " ".join(cmd), "findings": []}
    try:
        rep["validator"] = json.loads(cp.stdout)
    except Exception:
        rep["validator_raw"] = (cp.stdout + cp.stderr)[-1500:]
    if cp.returncode != 0:
        rep["findings"].append({"severity": "FAIL", "check": "validate_dataset", "detail": f"exit {cp.returncode}; failing checks are in validator.checks"})
    rep["fail"] = cp.returncode != 0
    return finish(rep)


def cmd_hebrew(args):
    rep = {"kind": "hebrew", "file": os.path.abspath(args.file), "findings": []}
    try:
        lines = extract_text(args.file)
    except Exception as e:
        rep["findings"].append({"severity": "FAIL", "check": "text-extract", "detail": f"{type(e).__name__}: {e}"})
        rep["fail"] = True
        return finish(rep)
    hits = hebrew_scan(lines)
    rep["stats"] = {"lines": len(lines), "hebrew_lines": sum(1 for _, l in lines if HEBREW.search(l)), "hits": len(hits)}
    for h in hits[:80]:
        rep["findings"].append({"severity": "FAIL", "check": h["rule"], "detail": f"{h['where']}: {h['token']}"})
    if len(hits) > 80:
        rep["findings"].append({"severity": "WARN", "check": "truncated", "detail": f"{len(hits) - 80} more hits not listed"})
    rep["note"] = "Every hit is a candidate. Read the line: a Latin run that is a proper noun inside Hebrew prose, a bracket, or a table with Hebrew is a violation; an all-Latin line, a code block, or a hyphenated number is not."
    rep["fail"] = bool(hits)
    return finish(rep)


# ----------------------------------------------------------------------------- main

def main(argv):
    if len(argv) >= 3 and argv[0] == "--worker":
        kind, payload = argv[1], json.loads(argv[2])
        print(json.dumps(_run_com(WORKERS[kind], payload), ensure_ascii=False))
        return 0
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("kind", choices=["xlsx", "docx", "pptx", "pdf", "html", "dataset", "hebrew"])
    ap.add_argument("file")
    ap.add_argument("--out")
    ap.add_argument("--dpi", type=int, default=80)
    ap.add_argument("--max-pages", type=int, default=12)
    ap.add_argument("--timeout", type=int, default=90)
    ap.add_argument("--json", action="store_true", help="accepted for symmetry; output is always JSON")
    ap.add_argument("--expect", help="xlsx: Sheet!A1=value,B2=value")
    ap.add_argument("--require-formulas", action="store_true")
    ap.add_argument("--rtl", action="store_true", help="docx: check Hebrew paragraphs are right-aligned in Word")
    ap.add_argument("--expect-a4", action="store_true")
    ap.add_argument("--expect-font")
    ap.add_argument("--viewport", default="1280x800")
    ap.add_argument("--mobile", action="store_true")
    args, passthrough = ap.parse_known_args(argv)
    if args.kind != "dataset" and passthrough:
        ap.error(f"unknown arguments: {passthrough}")
    if args.kind != "html" and not os.path.exists(args.file):
        print(json.dumps({"kind": args.kind, "file": args.file, "status": "FAIL", "findings": [{"severity": "FAIL", "check": "file-exists", "detail": "no such file"}]}))
        return 1
    return {"xlsx": cmd_xlsx, "docx": cmd_docx, "pptx": cmd_pptx, "pdf": cmd_pdf, "html": cmd_html,
            "dataset": lambda a: cmd_dataset(a, passthrough), "hebrew": cmd_hebrew}[args.kind](args)


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
